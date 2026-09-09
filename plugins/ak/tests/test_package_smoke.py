from __future__ import annotations

import base64
import json
import re
import subprocess
import sys
from pathlib import Path

import pytest
import yaml


PACKAGE = Path(__file__).resolve().parents[1]
REPOSITORY = PACKAGE.parents[1]
SCRIPTS = PACKAGE / "scripts"
sys.path.insert(0, str(SCRIPTS))



def run_script(name: str, *args: str) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, str(SCRIPTS / name), *map(str, args)],
        cwd=PACKAGE,
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert result.returncode == 0, f"{name} failed\nSTDOUT:\n{result.stdout}\nSTDERR:\n{result.stderr}"
    return result


def test_public_package_contract() -> None:
    import re

    package = json.loads((PACKAGE / "specifications" / "package.json").read_text(encoding="utf-8"))
    version = package["version"]
    assert re.fullmatch(r"\d+\.\d+\.\d+", version), version
    assert package["architecture_inspiration"]["dependency"] is False
    assert package["architecture_inspiration"]["vendored_code"] is False

    # package.json is the single source of truth; every manifest bump_version.py
    # touches must stay in lock-step so one release command is sufficient.
    for manifest in (".codex-plugin/plugin.json", ".claude-plugin/plugin.json"):
        assert json.loads((PACKAGE / manifest).read_text(encoding="utf-8"))["version"] == version
    marketplace = json.loads((REPOSITORY / ".claude-plugin" / "marketplace.json").read_text(encoding="utf-8"))
    assert marketplace["metadata"]["version"] == version
    assert all(entry["version"] == version for entry in marketplace["plugins"])

    run_script("validate_structure.py", "--package", str(PACKAGE), "--repository-root", str(REPOSITORY))


def test_synthetic_collaboration_fixture_has_no_conflicts() -> None:
    fixture = PACKAGE / "fixtures" / "collaboration" / "two-contributor"
    result = run_script(
        "ak.py",
        "collaboration",
        "package",
        "conflicts",
        "--root",
        str(fixture / "collaboration" / "work-packages"),
    )
    assert json.loads(result.stdout)["status"] == "VALID"


def test_public_yaml_has_unique_keys() -> None:
    class UniqueKeyLoader(yaml.SafeLoader):
        pass

    def construct_mapping(loader: UniqueKeyLoader, node: yaml.MappingNode, deep: bool = False) -> dict:
        mapping: dict = {}
        for key_node, value_node in node.value:
            key = loader.construct_object(key_node, deep=deep)
            assert key not in mapping, f"Duplicate YAML key {key!r} at line {key_node.start_mark.line + 1}"
            mapping[key] = loader.construct_object(value_node, deep=deep)
        return mapping

    UniqueKeyLoader.add_constructor(
        yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG,
        construct_mapping,
    )
    yaml_files = sorted((REPOSITORY / ".github").rglob("*.yml"))
    yaml_files += [REPOSITORY / "CITATION.cff", PACKAGE / "examples" / "minimal-app" / "manifest.yaml"]
    for path in yaml_files:
        yaml.load(path.read_text(encoding="utf-8"), Loader=UniqueKeyLoader)


def test_friendly_cli_entrypoint(tmp_path: Path) -> None:
    run_script("ak.py", "validate")
    run_script(
        "ak.py", "install",
        "--runtime", "generic",
        "--destination", str(tmp_path / "ak"),
        "--dry-run",
    )
    run_script(
        "ak.py", "install",
        "--runtime", "claude",
        "--project", str(tmp_path / "claude-project"),
        "--dry-run",
    )
    claude_project = tmp_path / "claude-project"
    run_script("ak.py", "install", "--runtime", "claude", "--project", str(claude_project))
    assert (claude_project / ".claude" / "ak-runtime").resolve() == PACKAGE.resolve()
    assert (claude_project / ".claude" / "skills" / "investigate").resolve() == (PACKAGE / "skills" / "investigate").resolve()
    run_script(
        "ak.py", "init",
        "--root", str(tmp_path),
        "--app-id", "T22",
        "--name-en", "Friendly CLI Test",
    )
    app = tmp_path / "T22"
    assert (app / "manifest.yaml").is_file()
    # Presentation output is optional and off by default.
    manifest_text = (app / "manifest.yaml").read_text(encoding="utf-8")
    assert "presentation_pptx: false" in manifest_text
    assert "graphify" not in manifest_text
    run_script("ak.py", "preflight", "--app-root", str(app))
    planned_documents = run_script(
        "ak.py", "documents", "--app-root", str(app), "--dry-run",
    )
    assert json.loads(planned_documents.stdout)["status"] == "PREFLIGHT_ONLY"


def test_synthetic_module_aware_pipeline(tmp_path: Path) -> None:
    run_script(
        "init_app.py",
        "--root", str(tmp_path),
        "--app-id", "T21",
        "--name-en", "Synthetic Public Test",
        "--runtime", "generic",
    )
    app = tmp_path / "T21"

    access_file = app / "input" / "access" / "synthetic.accdb"
    access_file.write_text("synthetic dry-run placeholder", encoding="utf-8")
    (app / "input" / "vba" / "DemoForm.bas").write_text(
        'Attribute VB_Name = "DemoForm"\nSub Save_Click(): End Sub\n', encoding="utf-8"
    )
    ignored = app / "input" / "sql" / "ignored.tmp"
    ignored.write_text("ignored", encoding="utf-8")

    compile_commands = app / "input" / "documents" / "compile_commands.json"
    compile_commands.write_text(
        json.dumps([{
            "directory": "/synthetic/build",
            "file": "/synthetic/demo.cpp",
            "arguments": ["clang++", "--password", "secret-value", "-c", "/synthetic/demo.cpp"],
        }]),
        encoding="utf-8",
    )

    session = app / "extracted" / "access" / "DB1" / "20260715-000000"
    session.mkdir(parents=True)
    (session / "component-index.json").write_text(
        json.dumps({
            "schema_version": "2.1",
            "app_id": "T21",
            "generated_at": "2026-07-15T00:00:00+00:00",
            "components": [{
                "id": "DB1:form:DemoForm",
                "kind": "form",
                "name": "DemoForm",
                "container": "demo",
                "module_hint": "demo",
                "source_paths": ["forms/DemoForm.txt"],
                "depends_on": [],
                "metadata": {},
            }],
        }),
        encoding="utf-8",
    )

    dry_run = run_script(
        "extract_access.py",
        "--database", str(access_file),
        "--database-id", "DB1",
        "--session-id", "DRY-RUN",
        "--output-dir", str(app / "extracted" / "access"),
        "--dry-run",
    )
    assert json.loads(dry_run.stdout)["status"] == "PREFLIGHT_ONLY"

    normalized = app / "extracted" / "build-context" / "compile_commands.normalized.json"
    run_script("parse_compilation_database.py", "--input", str(compile_commands), "--output", str(normalized))
    normalized_text = normalized.read_text(encoding="utf-8")
    assert "secret-value" not in normalized_text
    assert "<REDACTED>" in normalized_text

    run_script("build_component_index.py", "--app-root", str(app))
    run_script(
        "build_module_plan.py",
        "--component-index", str(app / ".ak" / "extracted" / "component-index.json"),
        "--output-dir", str(app / ".ak" / "extracted" / "module-plan"),
    )
    run_script("create_run.py", "--app-root", str(app), "--runtime", "generic", "--run-id", "T21-PUBLIC-SMOKE")
    run = app / ".ak" / "runs" / "T21-PUBLIC-SMOKE"
    run_script("create_tasks.py", "--package", str(PACKAGE), "--run", str(run))

    inventory = json.loads((run / "source-inventory.json").read_text(encoding="utf-8"))
    assert any(item["relative_path"].endswith("ignored.tmp") for item in inventory["ignored"])
    tasks = [json.loads(path.read_text(encoding="utf-8")) for path in sorted((run / "tasks").glob("*.json"))]
    assert tasks
    assert any(task["module_targets"] for task in tasks)
    # Derivation replaced six per-phase graph gates with one task, because the bundle
    # it reads is sealed and does not change between phases.
    derive_tasks = [task for task in tasks if task["role"] == "fact_deriver"]
    assert len(derive_tasks) == 1, "fact derivation runs once, not once per phase"


def test_access_runtime_probe_reports_json() -> None:
    result = run_script("access_runtime.py")
    report = json.loads(result.stdout)
    for key in ("platform", "python_process_bitness", "registry", "powershell_candidates", "selected_host", "activation", "status"):
        assert key in report, f"access runtime report missing {key}"
    assert report["activation"]["status"] in {"NOT_REQUESTED", "READY", "NOT_INSTALLED", "REGISTERED_BUT_ACTIVATION_FAILED", "INSTALLED_BUT_BITNESS_MISMATCH", "NOT_FOUND"}
    assert report["selected_host"]["status"] in {"READY", "NOT_INSTALLED", "INSTALLED_BUT_BITNESS_MISMATCH", "NOT_FOUND"}


def test_extract_access_reports_runtime_block(tmp_path: Path) -> None:
    database = tmp_path / "runtime.accdb"
    database.write_text("synthetic dry-run placeholder", encoding="utf-8")

    discovered = run_script(
        "extract_access.py",
        "--database", str(database),
        "--database-id", "DBR",
        "--session-id", "DRY-RUNTIME",
        "--output-dir", str(tmp_path / "extracted"),
        "--dry-run",
    )
    plan = json.loads(discovered.stdout)
    assert plan["status"] == "PREFLIGHT_ONLY"
    runtime = plan["runtime"]
    assert runtime["runtime_check"] == "COMPLETED"
    assert runtime["runtime_tested"] is False
    assert "host" in runtime and "status" in runtime["host"]

    skipped = run_script(
        "extract_access.py",
        "--database", str(database),
        "--database-id", "DBR",
        "--session-id", "DRY-SKIP",
        "--output-dir", str(tmp_path / "extracted"),
        "--dry-run",
        "--skip-runtime-check",
    )
    skipped_runtime = json.loads(skipped.stdout)["runtime"]
    assert skipped_runtime["runtime_check"] == "SKIPPED"
    assert skipped_runtime["runtime_tested"] is False


def test_preflight_input_preconditions(tmp_path: Path) -> None:
    run_script(
        "init_app.py",
        "--root", str(tmp_path),
        "--app-id", "T24",
        "--name-en", "Preconditions Test",
        "--runtime", "generic",
    )
    app = tmp_path / "T24"
    manifest = app / "manifest.yaml"

    empty = json.loads(run_script("preflight.py", "--package", str(PACKAGE), "--manifest", str(manifest)).stdout)
    precond = empty["input_preconditions"]
    assert precond["mode"] == "none"
    assert set(precond["recommended_missing"]) == {"input/vba", "input/sql"}

    (app / "input" / "vba" / "Form1.bas").write_text('Attribute VB_Name = "Form1"\n', encoding="utf-8")
    (app / "input" / "sql" / "schema.sql").write_text("CREATE TABLE t(id int);\n", encoding="utf-8")
    exported = json.loads(run_script("preflight.py", "--package", str(PACKAGE), "--manifest", str(manifest)).stdout)
    assert exported["input_preconditions"]["mode"] == "export"
    assert exported["input_preconditions"]["recommended_missing"] == []

    access_db = app / "input" / "access" / "T24.accdb"
    access_db.write_text("synthetic placeholder", encoding="utf-8")
    extract = json.loads(run_script("preflight.py", "--package", str(PACKAGE), "--manifest", str(manifest)).stdout)
    # VBA/SQL exports already present, so an unextracted Access binary makes it mixed.
    assert extract["input_preconditions"]["mode"] == "mixed"
    assert "runtime_status" in extract["input_preconditions"]


def test_nested_manifest_sources_drive_preflight_and_task_inputs(tmp_path: Path) -> None:
    run_script(
        "init_app.py",
        "--root", str(tmp_path),
        "--app-id", "T25",
        "--name-en", "Nested Source Test",
        "--runtime", "generic",
    )
    app = tmp_path / "T25"
    manifest = app / "manifest.yaml"
    text = manifest.read_text(encoding="utf-8")
    text = text.replace(
        'vba_exports: ["input/vba"]',
        'vba_exports:\n    - "input/T25_FRONTEND/vba"\n    - "input/T25_DATA/vba"',
    )
    text = text.replace('exported_paths: ["input/sql"]', "exported_paths: []")
    manifest.write_text(text, encoding="utf-8")
    frontend = app / "input" / "T25_FRONTEND" / "vba"
    data = app / "input" / "T25_DATA" / "vba"
    frontend.mkdir(parents=True)
    data.mkdir(parents=True)
    (frontend / "Form1.bas").write_text('Attribute VB_Name = "Form1"\n', encoding="utf-8")
    (data / "DataModule.bas").write_text('Attribute VB_Name = "DataModule"\n', encoding="utf-8")
    (app / "input" / "access" / "T25.accdb").write_text("synthetic placeholder", encoding="utf-8")

    report = json.loads(run_script("preflight.py", "--package", str(PACKAGE), "--manifest", str(manifest)).stdout)
    preconditions = report["input_preconditions"]
    assert preconditions["mode"] == "mixed"
    assert preconditions["present"]["vba"] is True
    assert preconditions["present"]["sql"] is False
    assert preconditions["recommended_missing"] == []
    assert preconditions["present"]["present_paths"]["vba"] == [
        "input/T25_FRONTEND/vba",
        "input/T25_DATA/vba",
    ]

    run_script("create_run.py", "--app-root", str(app), "--runtime", "generic", "--run-id", "T25-NESTED")
    run = app / ".ak" / "runs" / "T25-NESTED"
    run_script("create_tasks.py", "--package", str(PACKAGE), "--run", str(run))
    tasks = [json.loads(path.read_text(encoding="utf-8")) for path in sorted((run / "tasks").glob("*.json"))]
    vba_task = next(task for task in tasks if task["role"] == "vba_ui")
    sql_task = next(task for task in tasks if task["role"] == "sql_data")
    assert "../../input/T25_FRONTEND/vba" in vba_task["input_paths"]
    assert "../../input/T25_DATA/vba" in vba_task["input_paths"]
    assert "../../input/sql" not in sql_task["input_paths"]


def test_extract_ps1_declares_unique_safe_names() -> None:
    # Regression guard (runs everywhere): the safe-name function must derive a
    # deterministic hash from the original name so distinct non-ASCII objects do
    # not sanitize to the same filename and overwrite each other on disk.
    ps1 = (SCRIPTS / "extract_access.ps1").read_text(encoding="utf-8")
    assert "function Get-SafeName" in ps1
    assert "ComputeHash" in ps1


def test_extract_ps1_safe_names_keep_the_original_object_name() -> None:
    """Export filenames must carry the object's real name.

    An earlier version replaced every non-``[A-Za-z0-9_.-]`` character with an
    underscore, which turned a whole Japanese application into unreadable
    ``_______-84e869a1.txt`` files - and then needed a hash suffix to undo the
    collisions that sanitize had just created. This asserts the opposite policy,
    matching ``tools/ExportAccessObjects.bas``: keep the name, alter only what the
    filesystem forbids, and suffix a digest only when something had to change.
    """
    import shutil

    powershell = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell:
        import pytest

        pytest.skip("PowerShell not available on this platform")

    script = (SCRIPTS / "extract_access.ps1").as_posix()
    command = (
        "$ErrorActionPreference='Stop';"
        f"$c = Get-Content -Raw -LiteralPath '{script}';"
        "$m = [regex]::Matches($c, '(?ms)^function Get-(SafeName|NameDigest)\\(.*?^\\}');"
        "if ($m.Count -lt 2) { throw 'Get-SafeName/Get-NameDigest not found' };"
        "$m | ForEach-Object { Invoke-Expression $_.Value };"
        "$names = @('共通ルーチン','q受注データ','Form1','a/b','c:d','');"
        "$out = ($names | ForEach-Object { Get-SafeName $_ }) -join [char]10;"
        # Base64 so the answer crosses the pipe as ASCII. Written plainly, PowerShell
        # emits it in the console code page and Python decodes it in the host locale:
        # cp932 read the Japanese names fine, cp1252 on the CI runner raised
        # UnicodeDecodeError inside subprocess's reader thread, which surfaced as
        # `result.stdout is None` and a green suite everywhere the maintainer looked.
        # The names are what is under test; how a console renders them is not.
        "[Convert]::ToBase64String([System.Text.Encoding]::UTF8.GetBytes($out))"
    )
    result = subprocess.run(
        [powershell, "-NoProfile", "-NonInteractive", "-Command", command],
        check=False,
        capture_output=True,
    )
    assert result.returncode == 0, result.stderr.decode("utf-8", "replace")
    payload = base64.b64decode(result.stdout.strip()).decode("utf-8")
    safe = [line.strip() for line in payload.splitlines() if line.strip()]
    assert len(safe) == 6, safe
    assert len(set(safe)) == 6, f"safe names collide: {safe}"
    japanese, query, ascii_name, slash, colon, empty = safe
    # A name the filesystem accepts is used verbatim - no sanitize, no digest.
    assert japanese == "共通ルーチン", japanese
    assert query == "q受注データ", query
    assert ascii_name == "Form1", ascii_name
    # Only genuinely illegal characters are replaced, and then the digest keeps the
    # altered name distinct from any other name that sanitized to the same string.
    assert slash.startswith("a_b-") and colon.startswith("c_d-"), (slash, colon)
    assert empty.startswith("object-"), empty


def test_both_exporters_forbid_the_same_characters() -> None:
    """`evidence-layout.yaml`: the two routes must write the same container names.

    They did not. `extract_access.ps1` asked the running platform for its invalid set
    via `GetInvalidFileNameChars()`, which on Linux is only NUL and `/`, so one object
    named `c:d` became `c:d` under pwsh and `c_d-256d2ec0` under Windows PowerShell -
    the same function, two names, and CI red on ubuntu only. The .bas had always named
    its list outright. Nothing compared the two, which is what this does.
    """
    ps1 = (SCRIPTS / "extract_access.ps1").read_text(encoding="utf-8")
    bas = (PACKAGE / "tools" / "ExportAccessObjects.bas").read_text(encoding="utf-8")

    declared = re.search(r"""\$illegal = \[regex\]::Escape\(.*?\+ '(.*?)'\)""", ps1)
    assert declared, "extract_access.ps1 no longer declares its own invalid set"
    ps1_chars = set(declared.group(1))

    listed = re.search(r'bad = Array\((.*?)\)', bas)
    assert listed, "ExportAccessObjects.bas no longer lists its replaced characters"
    # A VBA literal for one double quote is four of them, so unwrap the outer pair
    # before unescaping the inner one - stripping every quote turns `""""` into "".
    bas_chars = set()
    for item in listed.group(1).split(","):
        item = item.strip()
        if item.startswith("vb"):
            continue
        assert item.startswith('"') and item.endswith('"'), item
        bas_chars.add(item[1:-1].replace('""', '"'))

    assert bas_chars == ps1_chars, (sorted(bas_chars), sorted(ps1_chars))
    # The .bas names CR, LF and TAB separately; the .ps1 covers them with 0..31.
    assert "0..31" in ps1


def test_vba_export_tool_present() -> None:
    bas = (PACKAGE / "tools" / "ExportAccessObjects.bas").read_text(encoding="utf-8")
    assert "Public Sub ExportAccessObjects" in bas
    assert "SaveAsText" in bas
    # Keeps original names, only de-duplicates on real collision (no lossy sanitizing).
    assert "UniquePath" in bas
    # Module name must differ from the Sub name, or calling it errors with
    # "Expected variable or procedure, not module".
    assert 'Attribute VB_Name = "modExportAccess"' in bas
    # SaveAsText output (system codepage / Shift-JIS) is transcoded to UTF-8 so
    # every export file is one consistent encoding.
    assert "shift_jis" in bas and 'Charset = "UTF-8"' in bas
    # Each object is exported independently; failures are recorded, not fatal.
    assert "AddSkip" in bas
    # System/temp/ImportErrors tables are excluded from the schema export.
    assert "IsSystemOrJunkTable" in bas


def test_recommended_optional_evidence_template() -> None:
    tpl = (PACKAGE / "templates" / "recommended-optional-evidence.md").read_text(encoding="utf-8")
    assert "{{APP_ID}}" in tpl
    assert "optional, non-blocking" in tpl
    for phase in ("Phase 4", "Phase 5", "Phase 2"):
        assert phase in tpl
def test_a_declared_access_runtime_that_does_not_match_stops_the_run(tmp_path: Path) -> None:
    database = tmp_path / "app.mdb"
    database.write_bytes(b"synthetic-signature-only")
    completed = subprocess.run(
        [
            sys.executable, str(SCRIPTS / "extract_access.py"),
            "--database", str(database), "--database-id", "T99",
            "--output-dir", str(tmp_path / "out"), "--session-id", "s1", "--execute",
            "--access-path", str(tmp_path / "nowhere" / "MSACCESS.EXE"),
        ],
        check=False, capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    if "windows" not in sys.platform.lower():
        return
    report = json.loads(completed.stdout)
    declared = report["runtime"]["declared_runtime"]
    # Only meaningful where an Access COM server is registered to disagree with.
    if declared["matches"] is None or not declared["registered_paths"]:
        return
    assert declared["matches"] is False
    assert report["status"] == "BLOCKED"
    assert completed.returncode == 3
    assert any("Declared Access runtime" in warning for warning in report["warnings"])
    # The snapshot must not have been taken: the run stopped before touching the file.
    assert not (tmp_path / "out" / "T99" / "s1" / "snapshot").exists()


def test_adopt_existing_workspace_preserves_files(tmp_path: Path) -> None:
    app = tmp_path / "T23"
    original = app / "docs" / "scope.md"
    original.parent.mkdir(parents=True)
    original.write_text("preserve", encoding="utf-8")

    refused = subprocess.run(
        [
            sys.executable,
            str(SCRIPTS / "ak.py"),
            "init",
            "--app-root", str(app),
            "--app-id", "T23",
            "--name-en", "Adopted Existing App",
        ],
        cwd=PACKAGE,
        check=False,
        capture_output=True,
        text=True,
    )
    assert refused.returncode != 0
    assert not (app / "manifest.yaml").exists()

    run_script(
        "ak.py",
        "init",
        "--app-root", str(app),
        "--app-id", "T23",
        "--name-en", "Adopted Existing App",
        "--adopt-existing",
    )
    assert original.read_text(encoding="utf-8") == "preserve"
    assert (app / "manifest.yaml").is_file()

def test_extract_access_blocks_before_snapshot_when_runtime_unavailable(tmp_path: Path) -> None:
    database = tmp_path / "blocked.accdb"
    database.write_text("synthetic placeholder", encoding="utf-8")
    out_dir = tmp_path / "extracted"
    # A non-existent PowerShell host forces NOT_FOUND without ever activating Access.
    result = subprocess.run(
        [
            sys.executable, str(SCRIPTS / "extract_access.py"),
            "--database", str(database),
            "--database-id", "DBX",
            "--session-id", "S1",
            "--output-dir", str(out_dir),
            "--execute",
            "--powershell", str(tmp_path / "missing" / "powershell.exe"),
        ],
        cwd=PACKAGE,
        check=False,
        capture_output=True,
        text=True,
    )
    assert result.returncode == 3, f"STDOUT:\n{result.stdout}\nSTDERR:\n{result.stderr}"
    plan = json.loads(result.stdout)
    assert plan["status"] == "BLOCKED"
    # The block must happen before any snapshot copy.
    assert not (out_dir / "DBX" / "S1" / "snapshot").exists()


# Phase 5 requires DOCUMENT-class evidence, and a document the kit cannot read is a
# document the phase cannot cite. These readers moved out of the deleted Graphify
# corpus builder into normalize_documents.py precisely because they are Phase 5
# infrastructure and were never Graphify's.
def test_document_normalizers_cover_office_and_report_scanned_pdf(tmp_path: Path) -> None:
    import docx
    import openpyxl
    import pptx
    import pypdf

    run_script(
        "init_app.py",
        "--root", str(tmp_path),
        "--app-id", "T28",
        "--name-en", "Document Normalization Test",
    )
    app = tmp_path / "T28"
    documents = app / "input" / "documents"

    workbook = openpyxl.Workbook()
    workbook.active.title = "業務規則"
    workbook.active.append(["項目", "規則"])
    workbook.active.append(["受注", "必須"])
    workbook.save(documents / "rules.xlsx")

    word = docx.Document()
    word.add_heading("運用手順", level=1)
    word.add_paragraph("受注確認を実行する。")
    word.save(documents / "manual.docx")

    slides = pptx.Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "業務フロー"
    slide.placeholders[1].text = "入力から出力まで"
    slides.save(documents / "flow.pptx")

    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with (documents / "scan.pdf").open("wb") as handle:
        writer.write(handle)

    run_script("normalize_documents.py", "--app-root", str(app))
    audit = json.loads((app / ".ak" / "extracted" / "normalized" / "NORMALIZATION_AUDIT.json").read_text(encoding="utf-8"))
    statuses = {entry["source_path"]: entry["status"] for entry in audit["entries"]}
    assert statuses["input/documents/rules.xlsx"] == "NORMALIZED"
    assert statuses["input/documents/manual.docx"] == "NORMALIZED"
    assert statuses["input/documents/flow.pptx"] == "NORMALIZED"
    # The page is genuinely blank, so which gap it becomes depends on the host, and
    # all three are the same statement: this source reached the corpus as a gap rather
    # than being silently skipped. `OCR_REQUIRED` on a machine with no Tesseract,
    # `OCR_FAILED` if it errors, and `OCR_NO_TEXT` where Tesseract is installed and
    # correctly finds nothing on an empty page. Asserting only the first two made this
    # pass for the wrong reason - it was reading the absence of an OCR engine.
    assert statuses["input/documents/scan.pdf"] in {
        "OCR_REQUIRED", "OCR_FAILED", "OCR_NO_TEXT",
    }


# --- OCR: three ways a supplied image reported success and delivered nothing -------

def test_tesseract_is_found_off_path_and_an_explicit_path_wins(tmp_path, monkeypatch):
    """The installer every Windows instruction points at does not amend PATH.

    So `install Tesseract` ended with a working executable the kit could not see, and
    OCR_REQUIRED told the person who had just installed it that they had not. Measured
    on this project's own host: Tesseract 5.4.0 present, absent from PATH.
    """
    import normalize_documents as nd

    monkeypatch.delenv("AK_TESSERACT", raising=False)
    monkeypatch.setattr(nd.shutil, "which", lambda _name: None)

    installed = tmp_path / "Tesseract-OCR" / "tesseract.exe"
    installed.parent.mkdir(parents=True)
    installed.write_text("", encoding="utf-8")
    monkeypatch.setattr(nd, "_tesseract_fallbacks", lambda: (installed,))
    assert nd.find_tesseract() == str(installed)

    # An operator who names the path has answered the question, so a wrong name is a
    # different failure from a missing one and must not quietly fall through to a
    # different executable than the one they asked for.
    monkeypatch.setenv("AK_TESSERACT", str(tmp_path / "nowhere.exe"))
    assert nd.find_tesseract() is None
    monkeypatch.setenv("AK_TESSERACT", str(installed))
    assert nd.find_tesseract() == str(installed)

    monkeypatch.delenv("AK_TESSERACT", raising=False)
    monkeypatch.setattr(nd, "_tesseract_fallbacks", lambda: ())
    assert nd.find_tesseract() is None


def _fake_tesseract(monkeypatch, module, stdout: str, languages: set[str]):
    monkeypatch.setattr(module, "find_tesseract", lambda: "tesseract")
    monkeypatch.setattr(module, "tesseract_languages", lambda _exe: languages)
    monkeypatch.setattr(module, "_prepared_for_ocr", lambda image, _work: (image, []))

    class _Result:
        returncode = 0
        stdout = ""
        stderr = ""

    def _run(*_args, **_kwargs):
        result = _Result()
        result.stdout = stdout
        return result

    monkeypatch.setattr(module.subprocess, "run", _run)


def test_english_only_ocr_of_a_japanese_corpus_says_so(tmp_path, monkeypatch):
    """Running is right. Saying nothing is not.

    An English source OCRs correctly with `eng`, so refusing would strand it. But this
    kit's whole target population is Japanese, and `eng` against Japanese does not
    fail - it returns confident nonsense, which then reads as evidence. The expensive
    outcome is not a refusal; it is an answer nobody knows to doubt.
    """
    import normalize_documents as nd

    _fake_tesseract(monkeypatch, nd, stdout="some text", languages={"eng", "osd"})
    _text, parser, warnings = nd.ocr_images([tmp_path / "page.png"])
    assert parser == "tesseract:eng"
    assert any(w.startswith("TESSERACT_NO_JPN") for w in warnings), warnings

    _fake_tesseract(monkeypatch, nd, stdout="some text", languages={"jpn", "eng"})
    _text, parser, warnings = nd.ocr_images([tmp_path / "page.png"])
    assert parser == "tesseract:jpn+eng"
    assert warnings == []


def test_ocr_that_returns_nothing_is_a_gap_not_a_normalization(tmp_path, monkeypatch):
    """Exit 0 and an empty string is what Tesseract gives for an unreadable image.

    The corpus recorded NORMALIZED, a parser and a hash for a file that contributed
    not one character, and the only way to notice was to open the corpus and find an
    empty section. `normalize_source`'s caller turns the prefix before the colon into
    the gap status, so this surfaces as OCR_NO_TEXT beside the other unread sources.
    """
    import normalize_documents as nd

    _fake_tesseract(monkeypatch, nd, stdout="   \n  ", languages={"jpn", "eng"})
    with pytest.raises(RuntimeError) as caught:
        nd.ocr_images([tmp_path / "blank.png"])
    assert str(caught.value).startswith("OCR_NO_TEXT:")


def test_a_supplied_image_is_stripped_of_alpha_before_ocr(tmp_path):
    """Measured on a real A06 screenshot: RGBA reads empty, dropped it reads.

    Its alpha is uniformly opaque, so removing it changes no pixel - confirmed by
    `ImageChops.difference` finding no bounding box between the two. This is not an
    image-quality adjustment; it is the difference between text and an empty string
    that the kit was recording as a successful normalization.

    An image with no alpha is handed through untouched, so nothing is re-encoded for
    the sake of it.
    """
    fitz = pytest.importorskip("fitz")
    import normalize_documents as nd

    source = tmp_path / "shot.png"
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 20), True)
    pixmap.clear_with(255)
    pixmap.save(source)
    assert fitz.Pixmap(str(source)).alpha, "the fixture must reproduce the real shape"

    prepared, warnings = nd._prepared_for_ocr(source, tmp_path)
    assert warnings == []
    assert prepared != source, "the original is never modified in place"
    assert not fitz.Pixmap(str(prepared)).alpha

    opaque = tmp_path / "opaque.png"
    plain = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 20), False)
    plain.clear_with(255)
    plain.save(opaque)
    assert nd._prepared_for_ocr(opaque, tmp_path) == (opaque, [])


def test_an_undecodable_diagnostic_does_not_take_down_the_run(tmp_path, monkeypatch):
    """Tesseract writes diagnostics in the host locale, which here is cp932.

    Strict utf-8 decoding raised inside subprocess's own reader thread, left
    `result.stderr` as None, and `None.strip()` then killed the whole `$ak documents`
    run with an AttributeError - so one unreadable image lost every other source in
    the workspace instead of being recorded as a gap.

    `adapters/managed_access` learned exactly this from a Japanese-Windows PowerShell
    and fixed it there. The lesson never reached the normalizer, which is why this
    asserts the None case as well as the decode: a reader thread can fail for reasons
    that are not encoding, and a diagnostic nobody can read is not a reason to lose
    the name of the file that produced it.
    """
    import normalize_documents as nd

    monkeypatch.setattr(nd, "find_tesseract", lambda: "tesseract")
    monkeypatch.setattr(nd, "tesseract_languages", lambda _exe: {"jpn", "eng"})
    monkeypatch.setattr(nd, "_prepared_for_ocr", lambda image, _work: (image, []))

    class _Result:
        returncode = 1
        stdout = ""
        stderr = None

    monkeypatch.setattr(nd.subprocess, "run", lambda *_a, **_k: _Result())
    with pytest.raises(RuntimeError) as caught:
        nd.ocr_images([tmp_path / "page.png"])
    message = str(caught.value)
    assert message.startswith("OCR_FAILED:")
    assert "exited 1" in message, message


def test_a_japanese_image_name_does_not_reach_tesseract_as_a_path(tmp_path):
    """Leptonica opens the path with the C runtime's narrow API.

    So a Japanese name in the temp path arrives mangled on a cp932 host and every read
    fails - reported as OCR_FAILED for a file that was perfectly readable. Four of a
    real A06 workspace's report exports (新商品一覧表.png and three more) failed this
    way within an hour of the alpha fix that introduced it.

    Named from a digest, which is the remedy this kit already uses twice:
    `extract_access.ps1` appends one to an altered filename, and bundle filenames come
    from a hash of the logical id. The original name stays in the audit entry, which
    is where it is read.
    """
    fitz = pytest.importorskip("fitz")
    import normalize_documents as nd

    source = tmp_path / "新商品一覧表.png"
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 20), True)
    pixmap.clear_with(255)
    pixmap.save(source)

    prepared, warnings = nd._prepared_for_ocr(source, tmp_path)
    assert warnings == []
    assert prepared != source
    assert prepared.name.isascii(), (
        f"the temp name must survive a narrow-API open, got {prepared.name!r}"
    )
    assert source.stem not in prepared.name

    # Two different names must not collide on one digest, and the same name must be
    # stable - a temp path that changed per run would be harmless here and confusing
    # in a log.
    other = tmp_path / "棚卸表.png"
    pixmap.save(other)
    assert nd._prepared_for_ocr(other, tmp_path)[0] != prepared
    assert nd._prepared_for_ocr(source, tmp_path)[0] == prepared
